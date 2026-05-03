"""Build pipeline presentation from extracted assets."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

ASSETS = '/shared/jingchl6/material/lc-research/test/ppt_assets'
OUT = '/shared/jingchl6/material/lc-research/test/pipeline_presentation.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

BLUE = RGBColor(0x1F, 0x47, 0x7A)
GREEN = RGBColor(0x2E, 0x86, 0x3E)
ORANGE = RGBColor(0xD9, 0x7A, 0x1F)
GRAY = RGBColor(0x55, 0x55, 0x55)

def add_title(slide, text, subtitle=None, color=BLUE):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(1.0))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = color
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.5))
        sp = sub.text_frame.paragraphs[0]; sp.text = subtitle
        sp.font.size = Pt(18); sp.font.color.rgb = GRAY

def add_bullets(slide, items, left, top, width, height, size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame; tf.word_wrap = True
    for i, txt in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f'• {txt}'; p.font.size = Pt(size); p.font.color.rgb = GRAY; p.space_after = Pt(6)

def add_image(slide, path, left, top, width=None, height=None):
    if not os.path.exists(path):
        return None
    return slide.shapes.add_picture(path, left, top, width=width, height=height)

def add_caption(slide, text, left, top, width, color=GRAY, size=12, bold=False, align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.4))
    p = box.text_frame.paragraphs[0]; p.text = text; p.alignment = align
    p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold

# ------- SLIDE 1: Title -------
s = prs.slides.add_slide(prs.slide_layouts[6])
box = s.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(12.3), Inches(2.0))
p = box.text_frame.paragraphs[0]; p.text = 'Liquid-Cell TEM Analysis Pipeline'
p.alignment = PP_ALIGN.CENTER; p.font.size = Pt(48); p.font.bold = True; p.font.color.rgb = BLUE
sub = s.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(12.3), Inches(1.5))
p2 = sub.text_frame.paragraphs[0]; p2.text = 'Drift correction → Denoising → Segmentation → Classification'
p2.alignment = PP_ALIGN.CENTER; p2.font.size = Pt(24); p2.font.color.rgb = GRAY
p3 = sub.text_frame.add_paragraph(); p3.text = 'Per-particle analysis of nanoparticle dynamics'
p3.alignment = PP_ALIGN.CENTER; p3.font.size = Pt(18); p3.font.color.rgb = GRAY
p3.space_before = Pt(12)

# ------- SLIDE 2: Overview -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Pipeline Overview', 'Four-stage automatic workflow with a single user click')

# Draw flow boxes
stages = [
    ('1. Drift\nCorrection', 'User specifies\nparticle; track\nacross video', GREEN),
    ('2. Denoising', 'Self-supervised\nneural networks\n(UDVD / UMVD)', BLUE),
    ('3. Segmentation', 'SAM 3 mask\n→ tight particle\nROI per frame', ORANGE),
    ('4. Classification\n(future)', 'CNN → 3 classes\n(dataset still\ncollecting)', GRAY),
]
box_w = Inches(2.8); gap = Inches(0.3); start_x = Inches(0.5); y = Inches(3.0); h = Inches(2.3)
for i, (title, desc, color) in enumerate(stages):
    x = start_x + (box_w + gap) * i
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    tf = shape.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(18); p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph(); p2.text = desc; p2.font.size = Pt(12); p2.space_before = Pt(8)
    p2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); p2.alignment = PP_ALIGN.CENTER
    # Arrow between boxes
    if i < len(stages) - 1:
        arrow_x = x + box_w + Inches(0.02)
        arrow_y = y + h/2 - Inches(0.15)
        arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_x, arrow_y, Inches(0.25), Inches(0.3))
        arr.fill.solid(); arr.fill.fore_color.rgb = GRAY; arr.line.color.rgb = GRAY

add_bullets(s,
    ['Input: raw AVI from liquid-cell TEM (~1364 frames, 2048×2048)',
     'Output: denoised particle-only video + per-frame masks',
     'Only manual step: one click to mark the particle at frame 0'],
    Inches(0.5), Inches(5.6), Inches(12.3), Inches(1.7), size=14)

# ------- SLIDE 3: Stage 1 — Drift Correction (problem) -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Stage 1 · Drift Correction', 'Particle moves over time; we need per-frame position')
add_image(s, f'{ASSETS}/01_raw_frame0.png', Inches(0.5), Inches(1.6), height=Inches(5.3))
add_caption(s, 'Frame 0 (raw) — 2048×2048, noisy', Inches(0.5), Inches(6.95), Inches(5.3))
add_bullets(s, [
    'Stage drift + liquid motion cause the particle to wander',
    'Blur frames occur at transitions (stage shifts, refocus)',
    'We need a per-frame (x, y) position, auto-detected',
    'Must handle large jumps and periods of blur'
], Inches(7.0), Inches(2.0), Inches(6.0), Inches(4.5), size=16)

# ------- SLIDE 4: Stage 1 — Template matching approach -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Stage 1 · Template Matching (DM-inspired)', 'Reverse-engineered from Gatan ImageAlignment.dll')
add_image(s, f'{ASSETS}/02_template_overlay.png', Inches(0.5), Inches(1.6), height=Inches(5.3))
add_caption(s, 'Frame 0: 900×900 template (green) shifted to include substrate features', Inches(0.5), Inches(6.95), Inches(5.3))
add_bullets(s, [
    'One user click → particle center (1085, 880)',
    'Template shifted by (−200, −250) to capture context',
    'Bandpass filter (DoG) removes DC and noise, keeps structure',
    'Global NCC search at every frame',
    'Blur frames auto-detected by Laplacian-variance drop; skipped'
], Inches(7.0), Inches(2.0), Inches(6.0), Inches(4.5), size=16)

# ------- SLIDE 5: Stage 1 results -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Stage 1 · Tracking Results', '8 / 8 ground-truth frames correct after DM-style tuning')
positions = [('Frame 0', '03_tracking_f0.png'),
             ('Frame 89 (after 1st blur)', '03_tracking_f89.png'),
             ('Frame 624 (previously hard)', '03_tracking_f624.png')]
x = Inches(0.5); w = Inches(4.0); y = Inches(1.6)
for i, (label, fn) in enumerate(positions):
    add_image(s, f'{ASSETS}/{fn}', x + Inches(i*4.3), y, width=w)
    add_caption(s, label, x + Inches(i*4.3), Inches(5.8), w, bold=True, size=14)
add_bullets(s, [
    'Ablation: template size, offset, preprocessing all tested',
    'Best: 900×900 template + bandpass + shifted (−200,−250)',
    'Tracking preserves particle position through 6 blur transitions'
], Inches(0.5), Inches(6.3), Inches(12.3), Inches(1.0), size=14)

# ------- SLIDE 6: Stage 2 — Denoising problem -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Stage 2 · Denoising', 'Self-supervised: no clean reference needed')
add_image(s, f'{ASSETS}/04_cropped_raw.png', Inches(0.5), Inches(1.6), height=Inches(5.0))
add_caption(s, 'Raw tracked particle (noisy)', Inches(0.5), Inches(6.7), Inches(5.0))
add_bullets(s, [
    'Two blind-spot networks trained from the video itself:',
    '   • UDVD-MF (Science 2025) — 5-frame window',
    '   • UMVD (CVPRW 2024) — depth-wise separable U-Net',
    '',
    'Training input: tracked particle-only TIFF stack',
    'Why particle-only? Focuses capacity on the region of interest,',
    'reduces compute, avoids over-smoothing background clutter',
    '',
    'Future work: test generalizability — can one trained model',
    'denoise other videos of the same particle type?'
], Inches(6.5), Inches(1.7), Inches(6.5), Inches(5.5), size=14)

# ------- SLIDE 7: Stage 2 — Denoising results -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Stage 2 · Denoising — Raw | Denoised', 'Side-by-side sample frame')
# Stack 3 rows (UDVD1, UDVD2, UMVD)
y = Inches(1.5); h = Inches(1.75)
for i, (name, fn, score) in enumerate([
    ('UDVD pass-1',    '05_denoise_udvd1.png', 'PSNR 18.98, uPSNR 23.86'),
    ('UDVD pass-2',    '05_denoise_udvd2.png', 'PSNR 40.89, uPSNR 28.64'),
    ('UMVD (best)',    '05_denoise_umvd.png',  'val_loss 0.0145, visually best')]):
    row_y = y + Inches(i * 1.95)
    add_image(s, f'{ASSETS}/{fn}', Inches(0.5), row_y, height=h)
    tb = s.shapes.add_textbox(Inches(8.0), row_y + Inches(0.3), Inches(5.0), Inches(1.0))
    p = tb.text_frame.paragraphs[0]; p.text = name; p.font.size = Pt(20); p.font.bold = True; p.font.color.rgb = BLUE
    p2 = tb.text_frame.add_paragraph(); p2.text = score; p2.font.size = Pt(14); p2.font.color.rgb = GRAY

# ------- SLIDE 8: Stage 3 — Segmentation -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Stage 3 · Segmentation (SAM 3)', 'Per-frame particle mask using a central box prompt')
positions = [('UDVD1', '06_seg_udvd1.png'),
             ('UDVD2', '06_seg_udvd2.png'),
             ('UMVD',  '06_seg_umvd.png')]
for i, (label, fn) in enumerate(positions):
    x = Inches(0.5 + i*4.3)
    add_image(s, f'{ASSETS}/{fn}', x, Inches(1.6), width=Inches(4.0))
    add_caption(s, label, x, Inches(5.8), Inches(4.0), bold=True, size=14)
add_bullets(s, [
    'SAM 3 image model, box prompt at center (0.5, 0.5, 0.4, 0.4)',
    '100 % of frames produced a usable mask across all 3 denoised inputs',
    'Mask → bounding box → tight per-particle crop'
], Inches(0.5), Inches(6.3), Inches(12.3), Inches(1.0), size=14)

# ------- SLIDE 9: Stage 3 — Tight crop output -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Stage 3 · Tight Particle-only Output', 'Ready for downstream classification')
add_image(s, f'{ASSETS}/07_tight_crop.png', Inches(0.5), Inches(1.6), height=Inches(5.0))
add_caption(s, 'UMVD denoised, cropped to mask bbox (256×256 uniform)', Inches(0.5), Inches(6.7), Inches(5.0))
add_bullets(s, [
    'Each frame: mask → square bbox → crop denoised image',
    'Resize all crops to uniform 256×256 for downstream model input',
    '',
    'Output artifacts:',
    '   • cropped_umvd_tight.mp4 — visual',
    '   • cropped_umvd_tight.tif — training-ready stack',
    '',
    'The particle fills every frame — zoomed, aligned, denoised'
], Inches(6.5), Inches(1.7), Inches(6.5), Inches(5.5), size=14)

# ------- SLIDE 10: Stage 4 — Classification (future) -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Stage 4 · Classification (planned)', '3-class particle identification')
shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.8), Inches(12.3), Inches(1.0))
shape.fill.solid(); shape.fill.fore_color.rgb = ORANGE
tf = shape.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.text = 'Waiting on data — collection in progress'
p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
p.alignment = PP_ALIGN.CENTER

add_bullets(s, [
    'Input: tight-cropped particle stacks from Stage 3',
    'Model (planned): 2D / 3D CNN — e.g. ResNet-3D or video transformer',
    'Classes (3 expected):',
    '   • Class A / B / C — to be specified as labeled data arrives',
    '',
    'Training plan once data is ready:',
    '   • Class-balanced sampling',
    '   • Data augmentation (rotation, temporal flip, intensity jitter)',
    '   • Cross-validation on per-video basis',
    '',
    'Until then: this pipeline is ready to feed clean, tracked,',
    'particle-only stacks into whichever classifier is chosen'
], Inches(0.5), Inches(3.2), Inches(12.3), Inches(4.2), size=16)

# ------- SLIDE 11: Summary -------
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, 'Summary & Next Steps')
# Two-column layout
col_w = Inches(6.0); gap = Inches(0.3)

add_bullets(s, [
    'Automatic end-to-end pipeline with 1-click input',
    '',
    'Stage 1 — Drift correction: 8/8 GT frames correct',
    'Stage 2 — Denoising: UMVD visually best; UDVD w/ 2nd',
    '                pass gives highest PSNR',
    'Stage 3 — Segmentation: 100 % frame retention, SAM 3',
    '',
    'All per-stage outputs saved:',
    '   • tracked_ts900_overlay.mp4 (tracking)',
    '   • umvd_ts900_comparison.mp4 (denoising)',
    '   • seg_umvd_overlay.mp4 (segmentation)',
    '   • cropped_umvd_tight.mp4 (final crop)'
], Inches(0.5), Inches(1.5), col_w, Inches(5.5), size=14)

add_bullets(s, [
    'Next steps',
    '',
    '  1. Collect classification training data',
    '       (~100–500 samples per class)',
    '',
    '  2. Test denoiser generalizability',
    '       Apply trained UMVD to a new video',
    '       of the same particle type',
    '',
    '  3. Train classifier (Stage 4)',
    '',
    '  4. End-to-end deployment:',
    '       input AVI → output particle class labels',
    ''
], Inches(6.8), Inches(1.5), col_w, Inches(5.5), size=14)

prs.save(OUT)
print(f'Saved: {OUT}')
print(f'Slides: {len(prs.slides)}')
